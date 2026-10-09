from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import re
from fastapi import HTTPException

@dataclass
class PageText:
    page: int
    text: str
    heading: str = ''
    location_type: str = 'page'

SUPPORTED = {'.pdf', '.docx', '.txt', '.md', '.png', '.jpg', '.jpeg'}


def parse_file(filename: str, content: bytes, ocr: bool = False) -> list[PageText]:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED:
        raise HTTPException(422, 'Unsupported format. Use PDF, DOCX, TXT, MD, PNG or JPG.')
    try:
        pages = []
        if extension == '.pdf':
            import pymupdf
            pymupdf.TOOLS.mupdf_display_errors(False)
            pymupdf.TOOLS.mupdf_display_warnings(False)
            with pymupdf.open(stream=content, filetype='pdf') as document:
                if document.is_encrypted:
                    raise HTTPException(422, 'Password-protected PDF cannot be parsed. Export an unlocked copy locally.')
                if len(document) > 500:
                    raise HTTPException(422, 'Limit: 500 pages per PDF. Split the document locally.')
                for number, page in enumerate(document, 1):
                    text = page.get_text()
                    if not text.strip() and ocr:
                        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5))
                        text = image_ocr(pixmap.tobytes('png'))
                    if text.strip():
                        pages.append(PageText(number, text, text.strip().splitlines()[0][:120]))
        elif extension == '.docx':
            import zipfile
            with zipfile.ZipFile(BytesIO(content)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 40 * 1024 * 1024:
                    raise HTTPException(422, 'DOCX expands beyond the safe parsing limit.')
            from docx import Document
            document = Document(BytesIO(content))
            section = 1
            heading = 'Document'
            text = []
            for paragraph in document.paragraphs:
                if paragraph.style.name.startswith('Heading') and text:
                    pages.append(PageText(section, '\n'.join(text), heading, 'section'))
                    section += 1
                    text = []
                if paragraph.style.name.startswith('Heading'):
                    heading = paragraph.text[:120]
                text.append(paragraph.text)
            for table in document.tables:
                for row in table.rows:
                    text.append(' | '.join(cell.text for cell in row.cells))
            if text:
                pages.append(PageText(section, '\n'.join(text), heading, 'section'))
        elif extension in {'.txt', '.md'}:
            text = content.decode('utf-8-sig')
            for number, page in enumerate(text.split('\f'), 1):
                heading = next((line.lstrip('# ').strip() for line in page.splitlines() if line.strip()), 'Document')[:120]
                if page.strip():
                    pages.append(PageText(number, page, heading))
        else:
            if not ocr:
                raise HTTPException(422, 'Image text needs the local OCR option (best effort).')
            pages = [PageText(1, image_ocr(content), 'Best-effort local OCR')]
        pages = [page for page in pages if page.text.strip()]
        if not pages:
            raise HTTPException(422, 'No selectable text found. Try the OCR option. OCR runs locally and is best effort.')
        if sum(len(page.text) for page in pages) > 2_000_000:
            raise HTTPException(422, 'Extracted text is too large. Split the document locally.')
        return pages
    except HTTPException:
        raise
    except UnicodeDecodeError:
        raise HTTPException(422, 'Text files must be UTF-8 encoded.') from None
    except Exception:
        raise HTTPException(422, 'This file could not be parsed locally. Check its format or export it as plain text.') from None


def image_ocr(content: bytes) -> str:
    """Tesseract stdin/stdout: never place a plaintext OCR image on disk."""
    from PIL import Image
    import subprocess
    try:
        with Image.open(BytesIO(content)) as image:
            if image.width * image.height > 20_000_000:
                raise HTTPException(422, 'Image is too large for local OCR. Resize it below 20 megapixels.')
            output = BytesIO()
            image.convert('RGB').save(output, format='PNG')
        result = subprocess.run(['tesseract', 'stdin', 'stdout'], input=output.getvalue(),
                                stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=30, check=False)
        if result.returncode != 0:
            raise HTTPException(422, 'Local OCR failed. Check Tesseract and try a clearer image.')
        return result.stdout.decode('utf-8', errors='replace')
    except FileNotFoundError:
        raise HTTPException(503, 'Local OCR needs Tesseract. Install it before the demo: sudo apt install tesseract-ocr') from None
    except subprocess.TimeoutExpired:
        raise HTTPException(422, 'Local OCR timed out. Try a smaller, clearer image.') from None


def chunk_pages(pages: list[PageText], filename: str, max_tokens: int = 400, overlap_tokens: int = 50) -> list[dict]:
    """Approximate token count by 4 characters/token. Exact slices preserve citations."""
    if max_tokens <= overlap_tokens or max_tokens < 50 or overlap_tokens < 0:
        raise ValueError('Invalid chunk size/overlap')
    result = []
    size = max_tokens * 4
    overlap = overlap_tokens * 4
    for page in pages:
        start = 0
        while start < len(page.text):
            end = min(len(page.text), start + size)
            if end < len(page.text):
                boundary = page.text.rfind(' ', start + size // 2, end)
                if boundary > start:
                    end = boundary
            text = page.text[start:end]
            headings = re.findall(r'^#{1,6}\s+(.+)$', page.text[:start + 1], re.M)
            if text.strip():
                result.append({'filename': filename, 'page': page.page, 'heading': headings[-1][:120] if headings else page.heading,
                               'location_type': page.location_type, 'text': text})
            if end == len(page.text):
                break
            start = max(start + 1, end - overlap)
    return result
