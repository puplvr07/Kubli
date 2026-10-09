import asyncio
import json
import pytest
from app.services.parsing import PageText, chunk_pages, parse_file
from app.services.rag import rank, answer, lookup_term, missing_items, NOT_COVERED
from app.services import llm
from app.services.store import Vault
from app.schemas import empty_record, RAGModelResponse
from app.routers.library import bounded_filename
from fastapi.testclient import TestClient
from app.main import app


def make_chunk(text='Document: allergies as reported.', **overrides):
    return dict(id='chunk1', document_id='doc1', text=text, filename='fever.txt', page=3, heading='Fever', tag='guidelines', embedding=[1., 0.], embedding_model=llm.EMBEDDING_PROFILE, **overrides)


def test_long_upload_filename_keeps_supported_extension():
    filename = '[CPG] ' + ('Clinical Practice Guidelines ' * 10) + '.pdf'
    bounded = bounded_filename(filename)
    assert len(bounded) == 160
    assert bounded.endswith('.pdf')


def test_chunk_pages_exact_slices_and_page_numbers():
    text = 'Chest pain notes ' * 150
    pages = [PageText(2, text, 'Chest pain'), PageText(7, 'Document: allergies as reported.', 'Fever')]
    chunks = chunk_pages(pages, 'sample.pdf', max_tokens=100, overlap_tokens=10)
    assert len(chunks) > 2
    assert {c['page'] for c in chunks} == {2, 7}
    assert all(c['text'] in (text if c['page'] == 2 else pages[1].text) for c in chunks)
    assert all(c['filename'] == 'sample.pdf' for c in chunks)
    assert chunks[-1]['heading'] == 'Fever'


def test_plain_text_and_docx_sections():
    assert [p.page for p in parse_file('notes.txt', b'page one\fpage two')] == [1, 2]
    from docx import Document
    from io import BytesIO
    document = Document(); document.add_heading('History', 1); document.add_paragraph('Past history stated.')
    document.add_heading('Exam', 1); document.add_paragraph('Exam findings stated.')
    output = BytesIO(); document.save(output)
    pages = parse_file('notes.docx', output.getvalue())
    assert pages[-1].location_type == 'section'
    assert pages[-1].heading == 'Exam'


def test_retrieval_embedding_first_and_scope_filter():
    fever = make_chunk('Fever checklist. Document: temperature in the stated unit.')
    chest = dict(make_chunk('Chest pain checklist. Document: blood pressure.'), id='chest', filename='chest.pdf', page=9, embedding=[0.,1.])
    result = rank([chest, fever], 'How should fever temperature be documented?', [1.,0.], ['guidelines'])
    assert result[0]['filename'] == 'fever.txt' and result[0]['page'] == 3
    # Zero keyword overlap is still eligible when the embedding points to the chunk.
    assert rank([fever], 'How do I capture thermal readings?', [1.,0.], ['guidelines'])[0]['id'] == 'chunk1'
    assert rank([fever], 'fever', [1.,0.], ['textbook']) == []


def test_grounded_answer_and_valid_not_covered(tmp_path, monkeypatch):
    vault = Vault(tmp_path); token = vault.unlock('strong password')
    vault.index = [make_chunk('Fever: document temperature in the stated unit.')]
    async def embedding(question): return [1., 0.]
    responses = [RAGModelResponse(status='answered', answer='Fever: document temperature in the stated unit.',
                                  source_ids=['S1']),
                 RAGModelResponse(status='not_covered', answer=NOT_COVERED, source_ids=[])]
    async def generated(prompt, schema_model, validator_extra=None, max_attempts=2, *, data):
        value = responses.pop(0)
        value = validator_extra(value) if validator_extra else value
        return llm.ValidatedGeneration(value=value, metadata={'done_reason': 'stop', 'eval_count': 40})
    monkeypatch.setattr(llm, 'embed_query', embedding); monkeypatch.setattr(llm, 'generate_validated', generated)
    result = asyncio.run(answer(vault, 'fever temperature', ['guidelines'], token))
    assert result['citations'][0]['snippet'] in vault.index[0]['text']
    assert result['citations'][0]['page'] == 3
    assert result['status'] == 'answered'
    uncovered = asyncio.run(answer(vault, 'galaxy orbital mechanics', ['guidelines'], token))
    assert uncovered['status'] == 'not_covered' and uncovered['answer'] == NOT_COVERED
    vault.close()


def test_term_lookup_returns_exact_definition_without_generation(tmp_path):
    vault = Vault(tmp_path); token = vault.unlock('strong password')
    source = ('Cardiovascular Conditions\nANSWERS\nby heart\nWhat Is a Heart Attack?\n'
              'A heart attack occurs when the blood flow that brings oxygen-rich blood to the heart muscle '
              'is severely reduced or cut off. This is due to a buildup of plaque in coronary arteries.')
    vault.index = [dict(make_chunk(source), heading='Heart attack')]
    result = lookup_term(vault, 'Heart attack', ['guidelines'], token)
    assert result['status'] == 'answered'
    assert result['answer'].startswith('A heart attack occurs when')
    assert result['answer'] in source
    assert result['citations'][0]['snippet'] == result['answer']
    vault.close()


def test_term_lookup_supports_short_abbreviations_and_scope(tmp_path):
    vault = Vault(tmp_path); token = vault.unlock('strong password')
    source = 'MI means myocardial infarction in this glossary. Verify the surrounding context.'
    vault.index = [dict(make_chunk(source), tag='textbook')]
    assert lookup_term(vault, 'MI', ['textbook'], token)['status'] == 'answered'
    assert lookup_term(vault, 'MI', ['notes'], token)['status'] == 'not_covered'
    assert lookup_term(vault, 'missing term', ['textbook'], token)['answer'] == NOT_COVERED
    vault.close()


def test_completeness_never_diagnoses():
    record = empty_record().model_dump()
    chunk = make_chunk('Fever example.\nDocument: allergies as reported.\nDocument: temperature in the stated unit.')
    flags = missing_items(record, [chunk])
    assert len(flags) == 2 and all(f['item'].startswith('Consider documenting: ') for f in flags)
    record['allergies'] = ['No known allergies']; record['vitals']['temp_c'] = 37
    assert not missing_items(record, [chunk])


def test_upload_review_encrypted_index_delete_and_reindex(tmp_path, monkeypatch):
    monkeypatch.setenv('WARDNOTE_DATA_DIR', str(tmp_path))
    async def embedding(texts): return [[1., 0.] for _ in texts]
    monkeypatch.setattr(llm, 'embed', embedding)
    with TestClient(app) as client:
        token = client.post('/api/unlock', json={'password': 'strong password'}).json()['token']
        headers = {'Authorization': 'Bearer ' + token}
        content = b'Fever notes. Patient name: Maria Santos.\nContact: clinic@example.com.\nDocument: allergies as reported.'
        files = [('files', ('private.txt', content, 'text/plain'))]
        preflight = client.post('/api/library/upload', headers=headers, files=files)
        assert preflight.status_code == 200
        assert client.get('/api/library', headers=headers).json() == []
        review = preflight.json()['files'][0]
        assert review['flags']
        assert review['upload_name'] == 'private.txt'
        assert client.post('/api/library/upload', headers=headers, files=files, data={'proceed':'true'}).status_code == 409
        decisions = {flag['id']: ('remove' if flag['kind'] == 'Possible patient name' else 'keep') for flag in review['flags']}
        result = client.post('/api/library/upload', headers=headers, files=files, data={
            'proceed':'true', 'deid_decisions':json.dumps({'private.txt': decisions})
        })
        assert result.status_code == 200
        document_id = result.json()['documents'][0]['id']
        assert client.get('/api/status', headers=headers).json()['chunks'] == 1
        indexed_text = app.state.vault.index[0]['text']
        assert 'Maria Santos' not in indexed_text
        assert '[removed]' in indexed_text
        assert 'clinic@example.com' in indexed_text
        assert b'Maria Santos' not in (tmp_path/'vault.sqlite3').read_bytes()
        assert client.post('/api/library/reindex', headers=headers).json()['chunks_indexed'] == 1
        assert client.delete(f'/api/library/{document_id}', headers=headers).status_code == 200
        assert client.get('/api/status', headers=headers).json()['chunks'] == 0
        assert client.get('/api/library', headers=headers).json() == []


def test_reindex_does_not_resurrect_deleted_chunks(tmp_path, monkeypatch):
    monkeypatch.setenv('WARDNOTE_DATA_DIR', str(tmp_path))
    async def change_index_during_embedding(texts):
        app.state.vault.index.clear()
        return [[1., 0.] for _ in texts]
    monkeypatch.setattr(llm, 'embed', change_index_during_embedding)
    with TestClient(app) as client:
        token = client.post('/api/unlock', json={'password': 'strong password'}).json()['token']
        app.state.vault.index = [make_chunk()]
        response = client.post('/api/library/reindex', headers={'Authorization':'Bearer '+token})
        assert response.status_code == 409
        assert app.state.vault.index == []
